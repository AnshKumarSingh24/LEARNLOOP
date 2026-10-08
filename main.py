from __future__ import annotations
import base64, hashlib, hmac, html, io, json, os, re, secrets, sqlite3, time, uuid, zipfile
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
try:
    from pypdf import PdfReader
except ImportError: PdfReader = None
try:
    from docx import Document as DocxDocument
except ImportError: DocxDocument = None
try:
    from pptx import Presentation
except ImportError: Presentation = None
try:
    from PIL import Image
except ImportError: Image = None
try:
    from openai import OpenAI
except ImportError: OpenAI = None
try:
    import chromadb
except ImportError: chromadb = None
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
except ImportError:
    SimpleDocTemplate = Paragraph = Spacer = getSampleStyleSheet = letter = None

ROOT=Path(__file__).resolve().parent.parent; DATA=ROOT/'data'; UPLOADS=DATA/'uploads'; DB_PATH=DATA/'learning.db'; CHROMA_DIR=DATA/'chroma'; UPLOADS.mkdir(parents=True,exist_ok=True)
app=FastAPI(title='Adaptive Learning & Artifact API',version='0.2.0'); app.add_middleware(CORSMiddleware,allow_origins=['*'],allow_methods=['*'],allow_headers=['*'])

# Keep the active root API aligned with the document service used by the UI.
CAPTURE_API_DIR = Path(__file__).resolve().parent / "capture_api"
sys.path.insert(0, str(CAPTURE_API_DIR))
from routers.document_questions import router as document_questions_router
from routers.artifacts import router as artifact_generator_router

app.include_router(document_questions_router)
app.include_router(artifact_generator_router)
DOCUMENT_UPLOADS = Path(__file__).resolve().parent / "uploads"
DOCUMENT_UPLOADS.mkdir(parents=True, exist_ok=True)

def vector_collection():
    if chromadb is None: return None
    return chromadb.PersistentClient(path=str(CHROMA_DIR)).get_or_create_collection('document_chunks')

def connect():
    c=sqlite3.connect(DB_PATH); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c

def init_db():
    with connect() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,email TEXT UNIQUE,password_hash TEXT,created_at TEXT);
        CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY,user_id TEXT,title TEXT,filename TEXT,media_type TEXT,content TEXT,created_at TEXT,FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE);
        CREATE TABLE IF NOT EXISTS chunks(id TEXT PRIMARY KEY,document_id TEXT,position INTEGER,content TEXT,FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE);
        CREATE TABLE IF NOT EXISTS artifacts(id TEXT PRIMARY KEY,document_id TEXT,kind TEXT,content TEXT,created_at TEXT,FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE);
        CREATE TABLE IF NOT EXISTS courses(id TEXT PRIMARY KEY,title TEXT,document_id TEXT,goal TEXT,daily_minutes INTEGER,level INTEGER,created_at TEXT,FOREIGN KEY(document_id) REFERENCES documents(id));
        CREATE TABLE IF NOT EXISTS concepts(id TEXT PRIMARY KEY,course_id TEXT,title TEXT,description TEXT,difficulty INTEGER,module_index INTEGER,lesson_index INTEGER,prerequisites TEXT,FOREIGN KEY(course_id) REFERENCES courses(id) ON DELETE CASCADE);
        CREATE TABLE IF NOT EXISTS mastery(user_id TEXT,concept_id TEXT,mastery REAL DEFAULT 0,confidence REAL DEFAULT 0,attempts INTEGER DEFAULT 0,correct_answers INTEGER DEFAULT 0,last_seen TEXT,next_review TEXT,repetitions INTEGER DEFAULT 0,interval_days REAL DEFAULT 0,easiness_factor REAL DEFAULT 2.5,PRIMARY KEY(user_id,concept_id),FOREIGN KEY(concept_id) REFERENCES concepts(id) ON DELETE CASCADE);
        ''')
        columns={row['name'] for row in c.execute('PRAGMA table_info(documents)').fetchall()}
        if 'user_id' not in columns: c.execute('ALTER TABLE documents ADD COLUMN user_id TEXT')
        mastery_columns={row['name'] for row in c.execute('PRAGMA table_info(mastery)').fetchall()}
        if 'repetitions' not in mastery_columns: c.execute('ALTER TABLE mastery ADD COLUMN repetitions INTEGER DEFAULT 0')
        if 'interval_days' not in mastery_columns: c.execute('ALTER TABLE mastery ADD COLUMN interval_days REAL DEFAULT 0')
        if 'easiness_factor' not in mastery_columns: c.execute('ALTER TABLE mastery ADD COLUMN easiness_factor REAL DEFAULT 2.5')
init_db()

AUTH_SECRET=os.getenv('AUTH_SECRET','dev-only-change-this-secret'); bearer=HTTPBearer(auto_error=False)

def hash_password(password):
    salt=secrets.token_bytes(16); digest=hashlib.pbkdf2_hmac('sha256',password.encode(),salt,210000); return base64.urlsafe_b64encode(salt+digest).decode()
def verify_password(password,encoded):
    raw=base64.urlsafe_b64decode(encoded.encode()); salt,digest=raw[:16],raw[16:]; return hmac.compare_digest(digest,hashlib.pbkdf2_hmac('sha256',password.encode(),salt,210000))
def make_token(user_id):
    payload=base64.urlsafe_b64encode(json.dumps({'sub':user_id,'exp':int(time.time())+86400}).encode()).decode().rstrip('='); sig=hmac.new(AUTH_SECRET.encode(),payload.encode(),hashlib.sha256).hexdigest(); return payload+'.'+sig
def current_user(credentials:Optional[HTTPAuthorizationCredentials]=Depends(bearer)):
    if not credentials: raise HTTPException(401,'Bearer token required')
    try:
        payload,sig=credentials.credentials.split('.',1); expected=hmac.new(AUTH_SECRET.encode(),payload.encode(),hashlib.sha256).hexdigest()
        data=json.loads(base64.urlsafe_b64decode(payload+'='*(-len(payload)%4)))
        if not hmac.compare_digest(sig,expected) or data['exp']<time.time(): raise ValueError
        with connect() as c: user=c.execute('SELECT * FROM users WHERE id=?',(data['sub'],)).fetchone()
        if not user: raise ValueError
        return user
    except Exception: raise HTTPException(401,'Invalid or expired token')

def now(): return datetime.now(timezone.utc)
def split_chunks(text,size=1400,overlap=180):
    clean=re.sub(r'\s+',' ',text).strip(); return [clean[i:i+size] for i in range(0,max(1,len(clean)),size-overlap)]

def extract(filename,data):
    name=filename.lower()
    if name.endswith('.pdf'):
        if PdfReader is None: raise HTTPException(500,'Install pypdf to ingest PDFs')
        p=UPLOADS/f'{uuid.uuid4()}_{filename}'; p.write_bytes(data)
        try: return '\n'.join(page.extract_text() or '' for page in PdfReader(str(p)).pages)
        finally: p.unlink(missing_ok=True)
    if name.endswith('.docx'):
        if DocxDocument is None: raise HTTPException(500,'Install python-docx to ingest DOCX files')
        p=UPLOADS/f'{uuid.uuid4()}_{filename}'; p.write_bytes(data)
        try: return '\n'.join(x.text for x in DocxDocument(str(p)).paragraphs)
        finally: p.unlink(missing_ok=True)
    if name.endswith('.pptx'):
        if Presentation is None: raise HTTPException(500,'Install python-pptx to ingest PPTX files')
        p=UPLOADS/f'{uuid.uuid4()}_{filename}'; p.write_bytes(data)
        try:
            return '\n'.join(shape.text for slide in Presentation(str(p)).slides for shape in slide.shapes if hasattr(shape,'text'))
        finally: p.unlink(missing_ok=True)
    if name.endswith(('.png','.jpg','.jpeg','.webp')):
        if Image is None: return '[Image uploaded. Configure a vision model/OCR adapter to extract text.]'
        return f'[Image uploaded: {filename}; dimensions={Image.open(__import__("io").BytesIO(data)).size}]'
    return data.decode('utf-8',errors='ignore')

def fallback_concepts(text,level=1):
    ss=[x.strip() for x in re.split(r'(?<=[.!?])\s+',text) if len(x.strip())>35]; out=[]
    for i,s in enumerate(ss[:16]):
        words=re.findall(r'[A-Za-z][A-Za-z-]+',s); title=' '.join(words[:7]).strip(' ,.:;')
        if title: out.append({'id':f'concept_{i+1}','title':title,'description':s,'difficulty':min(5,max(1,level+i//4)),'module_index':i//4+1,'lesson_index':i%4+1,'prerequisites':[f'concept_{i}'] if i else []})
    return out or [{'id':'concept_1','title':'Core ideas','description':text[:400],'difficulty':level,'module_index':1,'lesson_index':1,'prerequisites':[]}]

def llm_json(prompt, schema_hint):
    if not (os.getenv('OPENAI_API_KEY') and OpenAI): return None
    r=OpenAI().chat.completions.create(model=os.getenv('LLM_MODEL','gpt-5-mini'),messages=[{'role':'system','content':'You are an evidence-grounded educational content designer. Use only the supplied source. Return valid JSON matching the requested shape.'},{'role':'user','content':prompt+'\nJSON shape:\n'+schema_hint}],response_format={'type':'json_object'},max_completion_tokens=6000)
    return json.loads(r.choices[0].message.content)

def lexical_retrieve(document_id,query,k=5):
    terms=set(re.findall(r'[a-z]{3,}',query.lower()))
    with connect() as c: rows=c.execute('SELECT * FROM chunks WHERE document_id=?',(document_id,)).fetchall()
    scored=[]
    for r in rows:
        words=set(re.findall(r'[a-z]{3,}',r['content'].lower())); scored.append((len(terms&words),r['position'],r['content']))
    return [x[2] for x in sorted(scored,key=lambda x:(x[0],-x[1]),reverse=True)[:k]]

def index_chunks(document_id, parts, user_id=None):
    collection=vector_collection()
    if collection is None: return False
    if parts:
        collection.upsert(ids=[f'{document_id}:{i}' for i in range(len(parts))],documents=parts,metadatas=[{'document_id':document_id,'user_id':user_id or 'legacy','position':i} for i in range(len(parts))])
    return True

def semantic_retrieve(document_id,query,k=5):
    collection=vector_collection()
    if collection is None: return []
    result=collection.query(query_texts=[query],n_results=k,where={'document_id':document_id})
    return result.get('documents',[[]])[0] if result.get('documents') else []

def library_retrieve(user_id,query,k=8):
    collection=vector_collection()
    if collection is not None:
        result=collection.query(query_texts=[query],n_results=k,where={'user_id':user_id})
        docs=result.get('documents',[[]])[0] if result.get('documents') else []
        metas=result.get('metadatas',[[]])[0] if result.get('metadatas') else []
        return [{'content':d,'document_id':m.get('document_id'),'position':m.get('position')} for d,m in zip(docs,metas)]
    with connect() as c: docs=c.execute('SELECT id FROM documents WHERE user_id=?',(user_id,)).fetchall()
    scored=[]
    for d in docs:
        for content in lexical_retrieve(d['id'],query,k): scored.append({'content':content,'document_id':d['id']})
    return scored[:k]

def owned_document(document_id,user_id):
    with connect() as c: d=c.execute('SELECT * FROM documents WHERE id=? AND user_id=?',(document_id,user_id)).fetchone()
    if not d: raise HTTPException(404,'Document not found in your library')
    return d

def retrieve(document_id,query,k=5):
    semantic=semantic_retrieve(document_id,query,k)
    return semantic or lexical_retrieve(document_id,query,k)

def save_document(user_id,filename,media_type,text):
    did=str(uuid.uuid4()); title=Path(filename).stem.replace('_',' ').replace('-',' ').title(); parts=split_chunks(text)
    with connect() as c:
        c.execute('INSERT INTO documents(id,user_id,title,filename,media_type,content,created_at) VALUES(?,?,?,?,?,?,?)',(did,user_id,title,filename,media_type,text,now().isoformat()))
        c.executemany('INSERT INTO chunks VALUES(?,?,?,?)',[(str(uuid.uuid4()),did,i,x) for i,x in enumerate(parts)])
    vector_indexed=index_chunks(did,parts,user_id)
    return did,title,len(parts),vector_indexed

def document_json(did):
    with connect() as c:
        d=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone();
        if not d: raise HTTPException(404,'Document not found')
        a=c.execute('SELECT id,kind,created_at FROM artifacts WHERE document_id=? ORDER BY created_at DESC',(did,)).fetchall()
    x=dict(d); x['artifacts']=[dict(v) for v in a]; return x

def course_json(cid,user_id=None):
    with connect() as c:
        course=c.execute('SELECT * FROM courses WHERE id=?',(cid,)).fetchone()
        if not course: raise HTTPException(404,'Course not found')
        rows=c.execute('SELECT * FROM concepts WHERE course_id=? ORDER BY module_index,lesson_index',(cid,)).fetchall(); result=dict(course); result['modules']=[]
        for r in rows:
            x=dict(r); x['prerequisites']=json.loads(x['prerequisites'])
            if user_id:
                m=c.execute('SELECT * FROM mastery WHERE user_id=? AND concept_id=?',(user_id,r['id'])).fetchone(); x['mastery']=dict(m) if m else {'mastery':0,'confidence':0}
            result['modules'].append(x)
        return result

from typing import Optional

class Progress(BaseModel):
    user_id:str=''; concept_id:str; correct:bool; confidence:float=Field(.7,ge=0,le=1); quality:Optional[int]=Field(None,ge=0,le=5)
class CourseRequest(BaseModel):
    user_id:str='demo-user'; goal:str='understand'; daily_minutes:int=Field(30,ge=5,le=240); level:int=Field(1,ge=1,le=5)
class ArtifactRequest(BaseModel):
    instruction:str='Create the most useful artifact for learning this material.'; user_id:str='demo-user'
class ChatRequest(BaseModel):
    user_id:str; question:str
class AuthRequest(BaseModel):
    email:str; password:str=Field(min_length=8)

@app.get('/health')
def health(): return {'status':'ok'}

@app.post('/auth/signup')
def signup(request:AuthRequest):
    uid=str(uuid.uuid4())
    try:
        with connect() as c: c.execute('INSERT INTO users VALUES(?,?,?,?)',(uid,request.email.lower().strip(),hash_password(request.password),now().isoformat()))
    except sqlite3.IntegrityError: raise HTTPException(409,'Email already registered')
    return {'access_token':make_token(uid),'token_type':'bearer','user':{'id':uid,'email':request.email.lower().strip()}}

@app.post('/auth/login')
def login(request:AuthRequest):
    with connect() as c: user=c.execute('SELECT * FROM users WHERE email=?',(request.email.lower().strip(),)).fetchone()
    if not user or not verify_password(request.password,user['password_hash']): raise HTTPException(401,'Invalid email or password')
    return {'access_token':make_token(user['id']),'token_type':'bearer','user':{'id':user['id'],'email':user['email']}}

@app.get('/auth/me')
def me(user=Depends(current_user)): return {'id':user['id'],'email':user['email']}

@app.post('/documents/upload')
async def upload(file:UploadFile=File(...),user=Depends(current_user)):
    data=await file.read(); text=extract(file.filename or 'upload.txt',data)
    if len(text.strip())<20: raise HTTPException(400,'Not enough extractable content')
    did,title,n,vector_indexed=save_document(user['id'],file.filename or 'upload.txt',file.content_type or 'text/plain',text)
    return {'document':document_json(did),'chunk_count':n,'vector_indexed':vector_indexed,'vector_backend':'chroma' if vector_indexed else 'lexical-fallback'}

@app.get('/documents/{document_id}')
def get_document(document_id:str,user=Depends(current_user)): owned_document(document_id,user['id']); return document_json(document_id)

@app.get('/documents/{document_id}/search')
def search_document(document_id:str,q:str,k:int=5,user=Depends(current_user)):
    owned_document(document_id,user['id'])
    results=semantic_retrieve(document_id,q,max(1,min(k,20)))
    backend='chroma'
    if not results:
        results=lexical_retrieve(document_id,q,max(1,min(k,20))); backend='lexical-fallback'
    return {'query':q,'backend':backend,'results':[{'rank':i+1,'content':x} for i,x in enumerate(results)]}

@app.get('/library/search')
def library_search(q:str,k:int=8,user=Depends(current_user)):
    results=library_retrieve(user['id'],q,max(1,min(k,30)))
    return {'query':q,'backend':'chroma' if chromadb is not None else 'lexical-fallback','results':[{'rank':i+1,**x} for i,x in enumerate(results)]}

ARTIFACTS={'summary','flashcards','quiz','mindmap','flowchart','notes','timeline','table','outline','study_guide'}

def fallback_artifact(kind,text,instruction):
    cs=fallback_concepts(text,1)
    if kind=='summary': return {'title':'Source summary','sections':[{'heading':x['title'],'content':x['description']} for x in cs[:8]]}
    if kind=='flashcards': return {'flashcards':[{'question':f'What is {x["title"]}?','answer':x['description'],'difficulty':'medium'} for x in cs[:12]]}
    if kind=='quiz': return {'questions':[{'question':f'Which statement best describes {x["title"]}?','answer':x['description'],'options':[x['description'],'Not covered by the source','A contradictory idea','An unrelated example']} for x in cs[:10]]}
    if kind in ('mindmap','outline'):
        return {'root':'Learning map','nodes':[{'id':'root','label':'Learning map','type':'root'}]+[{'id':x['id'],'label':x['title'],'type':'concept','description':x['description']} for x in cs],'edges':[{'from':'root','to':x['id']} for x in cs]}
    if kind=='flowchart': return {'nodes':[{'id':str(i),'label':x['title']} for i,x in enumerate(cs)],'edges':[{'from':str(i),'to':str(i+1)} for i in range(len(cs)-1)]}
    if kind=='timeline': return {'events':[{'label':x['title'],'description':x['description'],'order':i} for i,x in enumerate(cs)]}
    if kind=='table': return {'columns':['Concept','Evidence','Difficulty'],'rows':[[x['title'],x['description'],x['difficulty']] for x in cs]}
    if kind=='study_guide': return {'title':'Study guide','objectives':[x['title'] for x in cs[:8]],'key_points':[x['description'] for x in cs[:12]],'practice_prompts':[f'Explain {x["title"]} in your own words.' for x in cs[:8]]}
    return {'title':'Notes','bullets':[x['description'] for x in cs]}

def make_artifact(did,kind,instruction):
    if kind not in ARTIFACTS: raise HTTPException(400,f'Unsupported artifact. Choose one of: {sorted(ARTIFACTS)}')
    with connect() as c: d=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone()
    if not d: raise HTTPException(404,'Document not found')
    source='\n\n'.join(retrieve(did,instruction,8)) or d['content'][:12000]
    result=llm_json(f'Create a {kind} from this source. User instruction: {instruction}\nSOURCE:\n{source}',f'{{"artifact_type":"{kind}","content":object}}')
    content=result.get('content',result) if result else fallback_artifact(kind,source,instruction); aid=str(uuid.uuid4())
    with connect() as c: c.execute('INSERT INTO artifacts VALUES(?,?,?,?,?)',(aid,did,kind,json.dumps(content),now().isoformat()))
    return {'id':aid,'document_id':did,'kind':kind,'content':content,'grounded_chunks':len(source.split('\n\n'))}

@app.post('/documents/{document_id}/artifacts/{kind}')
def artifact(document_id:str,kind:str,request:ArtifactRequest,user=Depends(current_user)):
    owned_document(document_id,user['id']); return make_artifact(document_id,kind,request.instruction)

@app.get('/artifacts/{artifact_id}')
def get_artifact(artifact_id:str):
    with connect() as c: x=c.execute('SELECT * FROM artifacts WHERE id=?',(artifact_id,)).fetchone()
    if not x: raise HTTPException(404,'Artifact not found')
    out=dict(x); out['content']=json.loads(out['content']); return out

def pdf_bytes(title,lines):
    if SimpleDocTemplate is None: raise HTTPException(500,'Install reportlab to export PDF files')
    output=io.BytesIO(); doc=SimpleDocTemplate(output,pagesize=letter); styles=getSampleStyleSheet(); story=[Paragraph(html.escape(title),styles['Title']),Spacer(1,12)]
    for line in lines: story.append(Paragraph(html.escape(str(line)),styles['BodyText'])); story.append(Spacer(1,6))
    doc.build(story); return output.getvalue()

def course_lines(course):
    lines=[f"Goal: {course['goal']} | Daily minutes: {course['daily_minutes']} | Level: {course['level']}"]
    for i,m in enumerate(course['modules'],1): lines += [f"Module/Lesson {i}: {m['title']}",m['description'],f"Difficulty: {m['difficulty']}"]
    return lines

def scorm_zip(course):
    manifest='''<?xml version="1.0" encoding="UTF-8"?><manifest identifier="adaptive-course" version="1.2" xmlns="http://www.imsproject.org/xsd/imscp_rootv1p1p2" xmlns:adlcp="http://www.adlnet.org/xsd/adlcp_rootv1p2" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"><organizations default="org"><organization identifier="org"><title>COURSE_TITLE</title>LESSONS</organization></organizations><resources><resource identifier="res" type="webcontent" adlcp:scormtype="sco" href="index.html"><file href="index.html"/></resource></resources></manifest>'''
    lesson_html=''.join(f'<item identifier="lesson{i}" identifierref="res"><title>{html.escape(m["title"])}</title></item>' for i,m in enumerate(course['modules'],1)); body=''.join(f'<section><h2>{html.escape(m["title"])}</h2><p>{html.escape(m["description"])}</p></section>' for m in course['modules'])
    manifest=manifest.replace('COURSE_TITLE',html.escape(course['title'])).replace('LESSONS',lesson_html); page=f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(course["title"])}</title></head><body><h1>{html.escape(course["title"])}</h1>{body}<script>window.parent&&window.parent.API&&window.parent.API.LMSInitialize("");</script></body></html>'
    output=io.BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z: z.writestr('imsmanifest.xml',manifest); z.writestr('index.html',page)
    return output.getvalue()

@app.get('/courses/{course_id}/export/pdf')
def export_course_pdf(course_id:str,user=Depends(current_user)):
    with connect() as c:
        if not c.execute('SELECT id FROM courses WHERE id=? AND EXISTS (SELECT 1 FROM documents d WHERE d.id=courses.document_id AND d.user_id=?)',(course_id,user['id'])).fetchone(): raise HTTPException(404,'Course not found')
    course=course_json(course_id,user['id']); data=pdf_bytes(course['title'],course_lines(course)); return StreamingResponse(io.BytesIO(data),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{course_id}.pdf"'})

@app.get('/courses/{course_id}/export/scorm')
def export_course_scorm(course_id:str,user=Depends(current_user)):
    with connect() as c:
        if not c.execute('SELECT id FROM courses WHERE id=? AND EXISTS (SELECT 1 FROM documents d WHERE d.id=courses.document_id AND d.user_id=?)',(course_id,user['id'])).fetchone(): raise HTTPException(404,'Course not found')
    data=scorm_zip(course_json(course_id,user['id'])); return StreamingResponse(io.BytesIO(data),media_type='application/zip',headers={'Content-Disposition':f'attachment; filename="{course_id}-scorm.zip"'})

@app.get('/artifacts/{artifact_id}/export/pdf')
def export_artifact_pdf(artifact_id:str,user=Depends(current_user)):
    with connect() as c: x=c.execute('SELECT a.*,d.user_id,d.title FROM artifacts a JOIN documents d ON d.id=a.document_id WHERE a.id=? AND d.user_id=?',(artifact_id,user['id'])).fetchone()
    if not x: raise HTTPException(404,'Artifact not found')
    content=json.loads(x['content']); lines=[]
    if x['kind'] in ('mindmap','flowchart','outline'):
        lines += [f"Nodes: {len(content.get('nodes',[]))}",f"Connections: {len(content.get('edges',[]))}"]+[n.get('label','')+' — '+n.get('description','') for n in content.get('nodes',[])]
    else: lines=[json.dumps(content,indent=2)]
    data=pdf_bytes(f"{x['title']} — {x['kind']}",lines); return StreamingResponse(io.BytesIO(data),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{artifact_id}.pdf"'})

@app.post('/documents/{document_id}/create-course')
def create_course(document_id:str,request:CourseRequest,user=Depends(current_user)):
    d=owned_document(document_id,user['id']); request.user_id=user['id']
    source='\n\n'.join(retrieve(document_id,'important concepts prerequisites learning order',12)) or d['content'][:18000]
    items=llm_json(f'Create a personalized course. Goal={request.goal}; daily minutes={request.daily_minutes}; level={request.level}. Source:\n{source}','{"concepts":[{"title":string,"description":string,"difficulty":1,"module_index":1,"lesson_index":1,"prerequisites":[string]}]}')
    concepts=(items or {}).get('concepts') if items else None; concepts=concepts or fallback_concepts(source,request.level); cid=str(uuid.uuid4()); suffix=cid[:8]
    with connect() as c:
        c.execute('INSERT INTO courses VALUES(?,?,?,?,?,?,?)',(cid,d['title'],document_id,request.goal,request.daily_minutes,request.level,now().isoformat()))
        ids={x['title']:f'concept_{i+1}_{suffix}' for i,x in enumerate(concepts)}
        for i,x in enumerate(concepts):
            xid=ids.get(x['title'],f'concept_{i+1}_{suffix}'); prereq=[ids[p] for p in x.get('prerequisites',[]) if p in ids] if isinstance(x.get('prerequisites',[]),list) else []
            c.execute('INSERT INTO concepts VALUES(?,?,?,?,?,?,?,?)',(xid,cid,x['title'],x.get('description',''),int(x.get('difficulty',1)),int(x.get('module_index',i//4+1)),int(x.get('lesson_index',i%4+1)),json.dumps(prereq)))
    return {'course':course_json(cid,request.user_id),'personalization':{'goal':request.goal,'daily_minutes':request.daily_minutes,'level':request.level}}

@app.get('/courses/{course_id}')
def get_course(course_id:str,user_id:Optional[str]=None,user=Depends(current_user)):
    with connect() as c:
        if not c.execute('SELECT id FROM courses WHERE id=? AND EXISTS (SELECT 1 FROM documents d WHERE d.id=courses.document_id AND d.user_id=?)',(course_id,user['id'])).fetchone(): raise HTTPException(404,'Course not found')
    return course_json(course_id,user['id'])

def choose(user_id,cid):
    with connect() as c:
        cs=c.execute('SELECT * FROM concepts WHERE course_id=? ORDER BY module_index,lesson_index',(cid,)).fetchall(); ms={x['concept_id']:x for x in c.execute('SELECT * FROM mastery WHERE user_id=?',(user_id,)).fetchall()}; choices=[]
        for x in cs:
            m=ms.get(x['id']); mastery=m['mastery'] if m else 0; ps=json.loads(x['prerequisites'])
            if any(ms.get(p,{'mastery':0})['mastery']<.7 for p in ps): continue
            due=not m or not m['next_review'] or datetime.fromisoformat(m['next_review'])<=now(); score=(1-mastery)*.55+(1 if due else 0)*.25+x['difficulty']/5*.1+(.1 if not m else 0); choices.append((score,x,mastery))
        if not choices:return None
        _,x,m=max(choices,key=lambda z:z[0]); return {'concept_id':x['id'],'title':x['title'],'description':x['description'],'difficulty':x['difficulty'],'mastery':m,'lesson_steps':['Learn','Example','Quick check','Practice','Complete']}

@app.get('/courses/{course_id}/next')
def next_lesson(course_id:str,user_id:str='demo-user',user=Depends(current_user)):
    with connect() as c:
        if not c.execute('SELECT id FROM courses WHERE id=? AND EXISTS (SELECT 1 FROM documents d WHERE d.id=courses.document_id AND d.user_id=?)',(course_id,user['id'])).fetchone(): raise HTTPException(404,'Course not found')
    return choose(user['id'],course_id) or {'completed':True,'message':'No eligible lessons remain'}
@app.post('/progress')
def progress(p:Progress,user=Depends(current_user)):
    p.user_id=user['id']
    t=now()
    with connect() as c:
        if not c.execute('SELECT id FROM concepts WHERE id=?',(p.concept_id,)).fetchone(): raise HTTPException(404,'Concept not found')
        old=c.execute('SELECT * FROM mastery WHERE user_id=? AND concept_id=?',(p.user_id,p.concept_id)).fetchone(); attempts=(old['attempts'] if old else 0)+1; correct=(old['correct_answers'] if old else 0)+int(p.correct); quality=p.quality if p.quality is not None else (4 if p.correct else 2); ef=float(old['easiness_factor'] if old and old['easiness_factor'] else 2.5); reps=int(old['repetitions'] if old and old['repetitions'] else 0); interval=float(old['interval_days'] if old and old['interval_days'] else 0)
        ef=max(1.3,ef+(0.1-(5-quality)*(0.08+(5-quality)*0.02))); reps=0 if quality<3 else reps+1
        if quality<3: interval=1
        elif reps==1: interval=1
        elif reps==2: interval=6
        else: interval=round(interval*ef,2) if interval else 6
        mastery=max(0,min(1,(old['mastery'] if old else 0)*.7+(quality/5)*.3)); review=t+timedelta(days=interval)
        c.execute('INSERT INTO mastery(user_id,concept_id,mastery,confidence,attempts,correct_answers,last_seen,next_review,repetitions,interval_days,easiness_factor) VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(user_id,concept_id) DO UPDATE SET mastery=excluded.mastery,confidence=excluded.confidence,attempts=excluded.attempts,correct_answers=excluded.correct_answers,last_seen=excluded.last_seen,next_review=excluded.next_review,repetitions=excluded.repetitions,interval_days=excluded.interval_days,easiness_factor=excluded.easiness_factor',(p.user_id,p.concept_id,mastery,p.confidence,attempts,correct,t.isoformat(),review.isoformat(),reps,interval,ef))
    return {'mastery':round(mastery,3),'attempts':attempts,'quality':quality,'repetitions':reps,'interval_days':interval,'easiness_factor':round(ef,3),'next_review':review.isoformat()}

@app.post('/documents/{document_id}/chat')
def chat(document_id:str,request:ChatRequest,user=Depends(current_user)):
    owned_document(document_id,user['id'])
    context='\n\n'.join(retrieve(document_id,request.question,6));
    if not context: raise HTTPException(404,'Document not found or has no chunks')
    if not (os.getenv('OPENAI_API_KEY') and OpenAI): return {'answer':'Based on the source:\n'+context[:1800],'mode':'fallback','citations':context[:2]}
    r=OpenAI().chat.completions.create(model=os.getenv('LLM_MODEL','gpt-5-mini'),messages=[{'role':'system','content':'Answer only from the supplied source context and say when it is insufficient.'},{'role':'user','content':f'Context:\n{context}\n\nQuestion: {request.question}'}],max_completion_tokens=900)
    return {'answer':r.choices[0].message.content or 'No answer generated.','mode':'llm','citations':context[:2]}

@app.get('/courses/{course_id}/progress')
def course_progress(course_id:str,user_id:str='demo-user',user=Depends(current_user)):
    with connect() as c:
        if not c.execute('SELECT id FROM courses WHERE id=? AND EXISTS (SELECT 1 FROM documents d WHERE d.id=courses.document_id AND d.user_id=?)',(course_id,user['id'])).fetchone(): raise HTTPException(404,'Course not found')
    x=course_json(course_id,user['id']); vals=[m.get('mastery',{}).get('mastery',0) if isinstance(m.get('mastery'),dict) else m.get('mastery',0) for m in x['modules']]; return {'course_id':course_id,'overall_mastery':round(sum(vals)/len(vals),3) if vals else 0,'concepts':[{'id':m['id'],'title':m['title'],'mastery':m.get('mastery',{}).get('mastery',0) if isinstance(m.get('mastery'),dict) else m.get('mastery',0)} for m in x['modules']]}

if __name__=='__main__':
    import uvicorn; uvicorn.run('main:app',host='0.0.0.0',port=8080,reload=True)
